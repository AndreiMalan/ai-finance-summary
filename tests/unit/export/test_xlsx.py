"""Excel export: every sheet reads back, values only, formula injection neutralised, values unrounded."""
import io

import pytest
from openpyxl import load_workbook

from statbrief.export.bundle import ReportBundle
from statbrief.export.xlsx import FORMULA_TRIGGERS, SHEETS, build_xlsx


def _wb(bundle: ReportBundle):
    """Workbook built from the bundle and read back (formulas not evaluated)."""
    return load_workbook(io.BytesIO(build_xlsx(bundle)))


def _rows(ws) -> list[list]:
    """All rows of a sheet as value lists."""
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_every_sheet_reads_back_with_content():
    """The six sheets exist in order and carry the bundle's records."""
    b = ReportBundle.example()
    wb = _wb(b)
    assert wb.sheetnames == list(SHEETS)
    data = _rows(wb["Data"])
    assert data[0][:3] == ["Obs id", "Indicator", "Indicator label"] and len(data) == len(b.observations) + 1
    assert len(_rows(wb["Metrics"])) == len(b.metrics) + 1
    prov = _rows(wb["Provenance"])
    assert len(prov) == 1 + sum(len(c.observation_ids) for c in b.draft.claims)
    conflicts = _rows(wb["Conflicts"])
    assert {r[0] for r in conflicts[1:]} == {"value_conflict", "unit_mismatch"}
    groups = [r[0] for r in _rows(wb["Citations"])[1:]]
    assert groups == ["Uploaded files", "Public APIs", "Reference documents"]
    summary = {r[0]: r[1] for r in _rows(wb["Summary"]) if r and r[0]}
    assert summary["Approved by"] == "analyst@example.com" and summary["Approved at"].startswith("2025-01-20T09:30")


def test_no_cell_holds_a_formula_and_values_are_unrounded():
    """Values only (AC 48); numbers are written exactly as stored."""
    b = ReportBundle.example()
    wb = _wb(b)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate)
    pp = next(m for m in b.metrics if m.kind == "pp_change")
    values = {r[0]: r[6] for r in _rows(wb["Metrics"])[1:]}
    assert values[pp.metric_id] == pp.value  # -7.799999999999999, not rounded


@pytest.mark.parametrize("payload", ["=HYPERLINK(\"http://x\",\"y\")", "+1+1", "-2+3", "@SUM(A1)", "\t=1", "\r=1"])
def test_formula_injection_is_neutralised(payload):
    """Formula-like text from files or the model is stored as a quote-prefixed string, never a formula."""
    data = ReportBundle.example().model_dump()
    data["draft"]["title"] = payload
    data["draft"]["caveats"] = [payload]
    data["observations"][0]["indicator_label"] = payload
    wb = _wb(ReportBundle.model_validate(data))
    read_back = {payload, payload.replace(chr(13), chr(10))}  # XML parsing normalises a carriage return
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value in read_back]
    assert len(cells) >= 3
    assert all(c.data_type == "s" and c.quotePrefix for c in cells)
    assert payload.startswith(FORMULA_TRIGGERS)


def test_illegal_xml_characters_are_removed():
    """Control characters (which openpyxl refuses) are stripped instead of failing the export."""
    data = ReportBundle.example().model_dump()
    data["draft"]["caveats"] = ["bad\x01text\x1f"]
    rows = _rows(_wb(ReportBundle.model_validate(data))["Summary"])
    assert ["badtext"] in [r[:1] for r in rows]


def test_negative_numbers_stay_numbers():
    """Only text is quote-prefixed: a negative number is still a numeric cell."""
    wb = _wb(ReportBundle.example())
    cell = next(c for row in wb["Metrics"].iter_rows(min_row=2) for c in row if isinstance(c.value, float)
                and c.value < 0)
    assert cell.data_type == "n" and not cell.quotePrefix
