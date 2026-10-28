"""Excel export of an approved report (openpyxl, values only, AC 48).

Sheets: Summary (parameters, approval, findings, caveats, validation exclusions, gap fills) · Data (the
observations used, with provenance columns) · Metrics · Provenance (one row per claim × input observation)
· Conflicts (value conflicts with their resolution, and unit mismatches) · Citations (grouped).

Values only: no cell holds a formula. Text that a spreadsheet would read as a formula (starting with
= + - @, tab or carriage return) is stored as a string cell with Excel's quote prefix, so it shows
literally and is not evaluated when the cell is edited (formula injection). Characters that XML cannot
carry are removed. Numbers are written unrounded (rounding happens only in presentation, i.e. the PDF);
timestamps are written as ISO 8601 text because Excel datetimes cannot hold a timezone.
"""
import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from statbrief.analytics.combine import source_label
from statbrief.domain.observation import ApiProvenance, Observation
from statbrief.export.bundle import CITATION_GROUPS, ReportBundle
from statbrief.export.format import observation_locator

SHEETS = ("Summary", "Data", "Metrics", "Provenance", "Conflicts", "Citations")
# Leading characters that make spreadsheet software treat text as a formula (OWASP CSV/formula injection).
FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")
_LIST_SEP = "; "
_BOLD = Font(bold=True)
_GROUP_HEADINGS = dict(CITATION_GROUPS)


def _clean(value: Any) -> Any:
    """Cell-safe value: ISO text for datetimes, illegal XML characters removed from text, None kept."""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _put(ws: Worksheet, row: int, col: int, value: Any) -> None:
    """Write one value; formula-like text becomes a quote-prefixed string cell (never a formula)."""
    value = _clean(value)
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if isinstance(value, str) and value.startswith(FORMULA_TRIGGERS):
        cell.data_type = "s"
        cell.quotePrefix = True


class _Sheet:
    """A worksheet written row by row (blank rows included) through `_put`."""

    def __init__(self, ws: Worksheet):
        """Start writing at the first row."""
        self.ws, self.row = ws, 0

    def append(self, values: list[Any], bold: bool = False) -> None:
        """Write the next row (an empty list leaves a blank separator row); bold for headers."""
        self.row += 1
        for col, value in enumerate(values, start=1):
            _put(self.ws, self.row, col, value)
            if bold:
                self.ws.cell(row=self.row, column=col).font = _BOLD

    def table(self, header: list[str], rows: list[list[Any]]) -> None:
        """A bold header row followed by the rows."""
        self.append(header, bold=True)
        for r in rows:
            self.append(r)


def _summary(ws: _Sheet, b: ReportBundle) -> None:
    """Parameters, approval, findings, caveats, validation exclusions and gap fills."""
    p = b.params
    ws.table(["Field", "Value"], [
        ["Title", b.draft.title], ["Report id", b.report_id], ["Mode", p.mode],
        ["Period from", p.period_from], ["Period to", p.period_to], ["Regions", _LIST_SEP.join(p.regions)],
        ["Indicators", _LIST_SEP.join(p.indicators)], ["Approved by", b.approved_by],
        ["Approved at", b.approved_at], ["Generated at", b.generated_at],
    ])
    ws.append([])
    ws.table(["Executive summary", "Claim ids"],
           [[f.headline, _LIST_SEP.join(f.claim_ids)] for f in b.draft.executive_summary])
    ws.append([])
    ws.table(["Caveats"], [[c] for c in b.draft.caveats])
    ws.append([])
    ws.table(["Validation exclusion: file", "Worksheet", "Issue", "Column", "Rows", "Message"],
           [[e.filename, e.worksheet, e.issue_code, e.column, _LIST_SEP.join(map(str, e.rows)), e.message]
            for e in b.validation_exclusions])
    ws.append([])
    fill_rows: list[list[Any]] = []
    if b.gap_fill is not None:
        fill_rows = [[f.upload_indicator, f.indicator, f.region, f.period, "filled", f.obs_id, f.source_obs_id,
                      f"{f.api_name}: {f.series_id}"] for f in b.gap_fill.filled]
        fill_rows += [[None, g.indicator, g.region, g.period, f"not filled ({g.reason})", g.obs_id, None, None]
                      for g in b.gap_fill.not_filled]
    ws.table(["Gap fill: upload indicator", "Indicator", "Region", "Period", "Outcome", "Obs id",
                "Source obs id", "Source series"], fill_rows)


def _provenance_columns(o: Observation) -> list[Any]:
    """Source type, label, file id/name/sheet/row/column or API name/series/URL/retrieval time."""
    p = o.provenance
    if isinstance(p, ApiProvenance):
        return ["api", source_label(o), None, None, None, None, None, p.api_name, p.series_id, p.query_url,
                p.retrieved_at]
    return ["upload", source_label(o), p.file_id, p.filename, p.worksheet, p.row_number, p.column,
            None, None, None, None]


_PROVENANCE_HEADER = ["Source type", "Source", "File id", "Filename", "Worksheet", "Row", "Column", "API",
                      "Series id", "Query URL", "Retrieved at"]


def _data(ws: _Sheet, b: ReportBundle) -> None:
    """The observations used, unrounded, with provenance columns and keep-both labels."""
    ws.table(["Obs id", "Indicator", "Indicator label", "Region", "Region level", "Period", "Frequency",
                "Value", "Unit", "Unit multiplier", "Status flag", "Label (keep both)", *_PROVENANCE_HEADER],
           [[o.obs_id, o.indicator_code, o.indicator_label, o.region_code, o.region_level, o.period, o.frequency,
             o.value, o.unit, o.unit_multiplier, o.status_flag, b.labels.get(o.obs_id), *_provenance_columns(o)]
            for o in b.observations])


def _metrics(ws: _Sheet, b: ReportBundle) -> None:
    """Every computed metric with its exact value, formula and input observations."""
    ws.table(["Metric id", "Kind", "Indicator", "Region", "Period", "Base period", "Value", "Unit",
                "Unit multiplier", "Formula", "Input obs ids"],
           [[m.metric_id, m.kind, m.indicator, m.region, m.period, m.base_period, m.value, m.unit,
             m.unit_multiplier, m.formula, _LIST_SEP.join(m.observation_ids)] for m in b.metrics])


def _provenance(ws: _Sheet, b: ReportBundle) -> None:
    """One row per claim and input observation: claim, metric, observation value and where it comes from."""
    rows = []
    for c in b.draft.claims:
        for oid in c.observation_ids:
            o = b.observation(oid)
            if o is None:  # unreachable: ReportBundle validation guarantees every input is carried
                continue
            rows.append([c.claim_id, c.text, c.value, c.unit, c.metric_id, oid, o.value, o.unit,
                         observation_locator(o), *_provenance_columns(o)])
    ws.table(["Claim id", "Claim text", "Claim value", "Claim unit", "Metric id", "Obs id", "Obs value",
                "Obs unit", "Locator", *_PROVENANCE_HEADER], rows)


def _conflicts(ws: _Sheet, b: ReportBundle) -> None:
    """Value conflicts (one row per candidate, with the resolution) and unit mismatches."""
    excluded = {e.obs_id: e.reason for e in b.conflict_exclusions}
    rows: list[list[Any]] = []
    for note in b.conflicts:
        c, r = note.conflict, note.resolution
        for oid, value, src in zip(c.candidates, c.values, c.sources, strict=True):
            rows.append(["value_conflict", c.conflict_id, c.indicator, c.region, c.period, c.unit, c.unit_multiplier,
                         oid, value, src, r.choice, r.preferred_obs_id, c.rel_diff,
                         "excluded" if oid in excluded else "kept", excluded.get(oid)])
    for m in b.unit_mismatches:
        rows.append(["unit_mismatch", None, m.indicator, m.region, m.period, _LIST_SEP.join(m.units), None,
                     _LIST_SEP.join(m.candidates), None, None, None, None, None, "kept separately", None])
    ws.table(["Type", "Conflict id", "Indicator", "Region", "Period", "Unit(s)", "Unit multiplier",
                "Candidate obs id(s)", "Value", "Source", "Resolution", "Preferred obs id", "Relative difference",
                "Outcome", "Reason"], rows)


def _citations(ws: _Sheet, b: ReportBundle) -> None:
    """Citations in group order (Uploaded files, Public APIs, Reference documents)."""
    ws.table(["Group", "Kind", "Ref id", "Label", "Source", "Locator", "URL"],
           [[_GROUP_HEADINGS[c.kind], c.kind, c.ref_id, c.label, c.source_label, c.locator, c.url]
            for _, cites in b.citations_by_group() for c in cites])


def build_xlsx(bundle: ReportBundle) -> bytes:
    """Render the bundle as an .xlsx workbook with the SHEETS above (values only, injection-safe)."""
    wb = Workbook()
    wb.remove(wb.active)
    writers = {"Summary": _summary, "Data": _data, "Metrics": _metrics, "Provenance": _provenance,
               "Conflicts": _conflicts, "Citations": _citations}
    for name in SHEETS:
        writers[name](_Sheet(wb.create_sheet(name)), bundle)
    wb.properties.creator = "StatBrief"
    wb.properties.title = _clean(bundle.draft.title)[:255]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

