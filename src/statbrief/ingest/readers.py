"""Read uploaded bytes into a grid of cells (no type coercion, no formula evaluation), then into tables."""
import csv
import io
import json
import struct
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

import xlrd

from statbrief.domain.dataset import FileKind
from statbrief.ingest.sniff import UploadRejected

_ENCODINGS = ("utf-8-sig", "cp1250", "latin-1")  # Romanian sources commonly use cp1250
_DELIMITERS = ",;\t|"
_SNIFF_BYTES = 64 * 1024
_JSON_TABLE_KEYS = ("data", "records", "rows")
# xlrd 2.0.2 raises XLRDError with exactly this text when it meets a FILEPASS record.
_XLRD_ENCRYPTED_MESSAGE = "Workbook is encrypted"
# BIFF records scanned in the workbook-globals substream (xlrd drops non-worksheet sheet types
# without recording them, so macro sheets are found by reading BOUNDSHEET records directly).
_BIFF_RECORD_HEADER = struct.Struct("<HH")  # record type, body length
_BIFF_EOF = 0x000A
_BIFF_BOUNDSHEET = 0x0085
_BOUNDSHEET_TYPE_OFFSET = 5  # body: stream offset (4 bytes), visibility (1), sheet type (1), name
_MACRO_SHEET_TYPES = frozenset({0x01, 0x06})  # Excel 4.0 (XLM) macro sheet, Visual Basic module
_MIN_SCANNED_BIFF_VERSION = 50  # BIFF5+: BOUNDSHEET records carry the sheet type
_PASSWORD_MESSAGE = ("This workbook is password-protected or encrypted. "
                     "Remove the password in Excel and upload it again.")
_MACRO_MESSAGE = "Workbooks containing macros are not accepted. Save it as a plain .xlsx without macros."


class _FormulaNoValue:
    """Sentinel type: a formula cell whose cached value is missing (never computed)."""

    def __repr__(self) -> str:
        """Readable representation in test output."""
        return "FORMULA_NO_VALUE"


FORMULA_NO_VALUE = _FormulaNoValue()


@dataclass
class Grid:
    """Non-empty source rows with their original 1-based row numbers."""

    rows: list[list[Any]]
    row_numbers: list[int]
    worksheet: str | None = None
    meta: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class SheetInfo:
    """A worksheet name and whether it is hidden."""

    name: str
    hidden: bool


@dataclass
class RawTable:
    """Header + data rows of one table, with original row numbers."""

    columns: list[str]
    rows: list[list[Any]]
    row_numbers: list[int]
    worksheet: str | None = None
    meta: dict[str, Any] = field(default_factory=dict)

    def column_index(self, name: str) -> int:
        """Position of a column (ValueError if absent)."""
        return self.columns.index(name)


def _is_empty(v: Any) -> bool:
    """True for None and blank strings."""
    return v is None or (isinstance(v, str) and not v.strip())


def _append(grid: Grid, cells: list[Any], row_number: int, max_rows: int | None) -> None:
    """Add a non-empty row, enforcing the row limit."""
    if all(_is_empty(c) for c in cells):
        return
    if max_rows is not None and len(grid.rows) >= max_rows:
        raise UploadRejected("too_large", f"The file has more than {max_rows} rows.")
    grid.rows.append(cells)
    grid.row_numbers.append(row_number)


def _read_csv(data: bytes, max_rows: int | None) -> Grid:
    """Decode with a fallback chain, sniff the delimiter, keep cells as strings.

    Upload sniffing only probes the first bytes for NULs, so any later NUL characters are removed
    here (the csv module refuses them) and flagged in `meta["nul_removed"]`.
    """
    for enc in _ENCODINGS:
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 always decodes
        raise UploadRejected("malformed", "The file's text encoding is not supported.")
    meta: dict[str, Any] = {"encoding": enc}
    if "\x00" in text:
        text = text.replace("\x00", "")
        meta["nul_removed"] = True
    try:
        delim = csv.Sniffer().sniff(text[:_SNIFF_BYTES], delimiters=_DELIMITERS).delimiter
    except csv.Error:
        delim = ","
    meta["delimiter"] = delim
    grid = Grid(rows=[], row_numbers=[], meta=meta)
    try:
        for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delim), start=1):
            _append(grid, [c if c != "" else None for c in row], i, max_rows)
    except csv.Error as exc:  # e.g. a field larger than the csv module's limit
        raise UploadRejected("malformed", "This CSV file could not be read.") from exc
    return grid


def _xlsx_workbooks(data: bytes):
    """Open a workbook twice: cached values, and raw cells to recognise formulas."""
    from openpyxl import load_workbook

    try:
        values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        raw = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise UploadRejected("malformed", "This Excel file could not be read.") from exc
    return values, raw


def _pick_sheet(infos: list[SheetInfo], worksheet: str | None) -> str:
    """Requested sheet, else the first visible one."""
    if worksheet is not None:
        if worksheet not in {s.name for s in infos}:
            raise UploadRejected("malformed", f"Worksheet {worksheet!r} does not exist.")
        return worksheet
    visible = [s.name for s in infos if not s.hidden]
    if not visible:
        raise UploadRejected("malformed", "The workbook has no visible worksheet.")
    return visible[0]


def list_worksheets(data: bytes, kind: FileKind) -> list[SheetInfo]:
    """Worksheets of an Excel file in workbook order; empty list for CSV/JSON."""
    if kind == "xlsx":
        values, raw = _xlsx_workbooks(data)
        try:
            return [SheetInfo(ws.title, ws.sheet_state != "visible") for ws in values.worksheets]
        finally:
            values.close()
            raw.close()
    if kind == "xls":
        book = _open_xls(data)
        try:
            return _xls_sheet_infos(book)
        finally:
            book.release_resources()
    return []


def _read_xlsx(data: bytes, worksheet: str | None, max_rows: int | None) -> Grid:
    """Cached cell values; formula cells without a cache become FORMULA_NO_VALUE."""
    values, raw = _xlsx_workbooks(data)
    try:
        infos = [SheetInfo(ws.title, ws.sheet_state != "visible") for ws in values.worksheets]
        name = _pick_sheet(infos, worksheet)
        grid = Grid(rows=[], row_numbers=[], worksheet=name)
        raw_rows = raw[name].iter_rows(values_only=True)
        for i, row in enumerate(values[name].iter_rows(values_only=True), start=1):
            raw_row = next(raw_rows, ())
            cells = []
            for j, v in enumerate(row):
                r = raw_row[j] if j < len(raw_row) else None
                cells.append(FORMULA_NO_VALUE if v is None and isinstance(r, str) and r.startswith("=") else v)
            _append(grid, cells, i, max_rows)
        return grid
    finally:
        values.close()
        raw.close()


def _xls_has_macro_sheets(stream: bytes, start: int, end: int) -> bool:
    """True if a BOUNDSHEET record in the globals substream declares a macro or VB-module sheet.

    Walks raw BIFF records from `start` (the globals BOF) up to the globals EOF record or `end`;
    only record headers and the one-byte sheet type are inspected.
    """
    pos = start
    while pos + _BIFF_RECORD_HEADER.size <= end:
        rc, length = _BIFF_RECORD_HEADER.unpack_from(stream, pos)
        body_start = pos + _BIFF_RECORD_HEADER.size
        if rc == _BIFF_EOF:
            return False
        if (rc == _BIFF_BOUNDSHEET and length > _BOUNDSHEET_TYPE_OFFSET
                and body_start + length <= end
                and stream[body_start + _BOUNDSHEET_TYPE_OFFSET] in _MACRO_SHEET_TYPES):
            return True
        pos = body_start + length
    return False


def _open_xls(data: bytes) -> xlrd.book.Book:
    """Open a legacy .xls (cached values only), rejecting encrypted and macro-sheet workbooks.

    The caller must call `release_resources()` on the returned book when done.
    """
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
    except xlrd.XLRDError as exc:
        if str(exc) == _XLRD_ENCRYPTED_MESSAGE:
            raise UploadRejected("password_protected", _PASSWORD_MESSAGE) from exc
        raise UploadRejected("malformed", "This .xls file could not be read.") from exc
    except Exception as exc:  # xlrd raises assorted errors on corrupt input
        raise UploadRejected("malformed", "This .xls file could not be read.") from exc
    try:
        if book.biff_version < _MIN_SCANNED_BIFF_VERSION:
            # Pre-BIFF5 files carry no sheet types we can check; sniffing already requires OLE2,
            # which only BIFF5+ uses, so this is defence in depth.
            raise UploadRejected("malformed", "This .xls format is too old to be read safely.")
        if _xls_has_macro_sheets(book.mem, book.base, book.base + book.stream_len):
            raise UploadRejected("macro_enabled", _MACRO_MESSAGE)
    except UploadRejected:
        book.release_resources()
        raise
    return book


def _xls_sheet_infos(book: xlrd.book.Book) -> list[SheetInfo]:
    """Names and hidden flags of a legacy workbook's worksheets (0 = visible in BOUNDSHEET)."""
    return [SheetInfo(s.name, bool(s.visibility)) for s in book.sheets()]


def _read_xls(data: bytes, worksheet: str | None, max_rows: int | None) -> Grid:
    """Legacy Excel via xlrd; date cells converted to datetime."""
    book = _open_xls(data)
    try:
        name = _pick_sheet(_xls_sheet_infos(book), worksheet)
        sheet = book.sheet_by_name(name)
        grid = Grid(rows=[], row_numbers=[], worksheet=name)
        for r in range(sheet.nrows):
            cells = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        cells.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                    except (ValueError, OverflowError):  # out-of-range serial: keep the number
                        cells.append(cell.value)
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    cells.append(None)
                else:
                    cells.append(cell.value)
            _append(grid, cells, r + 1, max_rows)
        return grid
    finally:
        book.release_resources()


def _read_json(data: bytes, max_rows: int | None) -> Grid:
    """List of objects (optionally wrapped) → header row + value rows."""
    try:
        doc = json.loads(data.decode("utf-8-sig"))
    except (ValueError, RecursionError) as exc:  # sniffing normally catches these first
        raise UploadRejected("malformed", "This JSON file could not be parsed.") from exc
    if isinstance(doc, dict):
        doc = next((doc[k] for k in _JSON_TABLE_KEYS if isinstance(doc.get(k), list)), None)
    if not isinstance(doc, list) or not doc or not all(isinstance(r, dict) for r in doc):
        raise UploadRejected("malformed", "This JSON is not a table. Use a list of objects, "
                             'e.g. [{"date": "2024-01", "region": "RO", "value": 1.2}].')
    columns: list[str] = []
    seen: set[str] = set()
    for rec in doc:
        for k in rec:
            if k not in seen:
                seen.add(k)
                columns.append(k)
    grid = Grid(rows=[], row_numbers=[], meta={})
    _append(grid, list(columns), 1, max_rows)
    for i, rec in enumerate(doc, start=2):
        cells = [rec.get(c) for c in columns]
        _append(grid, [json.dumps(v) if isinstance(v, (dict, list)) else v for v in cells], i, max_rows)
    return grid


def read_grid(data: bytes, kind: FileKind, worksheet: str | None = None, max_rows: int | None = None) -> Grid:
    """Read any supported file into a grid of raw cells."""
    if kind == "csv":
        return _read_csv(data, max_rows)
    if kind == "xlsx":
        return _read_xlsx(data, worksheet, max_rows)
    if kind == "xls":
        return _read_xls(data, worksheet, max_rows)
    return _read_json(data, max_rows)


def _header_name(value: Any, position: int) -> str:
    """Readable column name from a header cell (stripped, to match ColumnMapping)."""
    if _is_empty(value) or value is FORMULA_NO_VALUE:
        return f"column_{position}"
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def _unique_names(bases: list[str]) -> list[str]:
    """First occurrence keeps its name; repeats get `_2`, `_3`, ... skipping any name already taken."""
    reserved, names = set(bases), []
    used: set[str] = set()
    for base in bases:
        name, n = base, 1
        while name in used or (n > 1 and name in reserved):
            n += 1
            name = f"{base}_{n}"
        used.add(name)
        names.append(name)
    return names


def to_table(grid: Grid, header_row: int) -> RawTable:
    """Use the grid row with original number `header_row` as the header; later rows are data."""
    if header_row not in grid.row_numbers:
        raise ValueError(f"row {header_row} is empty or outside the table")
    h = grid.row_numbers.index(header_row)
    width = max([len(grid.rows[h])] + [len(r) for r in grid.rows[h + 1:]])
    header = grid.rows[h] + [None] * (width - len(grid.rows[h]))  # extra data columns get column_N
    names = _unique_names([_header_name(cell, pos) for pos, cell in enumerate(header, start=1)])
    rows = [r + [None] * (width - len(r)) for r in grid.rows[h + 1:]]
    return RawTable(columns=names, rows=rows, row_numbers=grid.row_numbers[h + 1:],
                    worksheet=grid.worksheet, meta=dict(grid.meta))
